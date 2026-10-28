"""Local XLSX adapter for the existing Census import pipeline. No external requests."""

import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert

from app.analytics.acs_profiles import PROFILE_INDICATORS
from app.core.config import ROOT
from app.core.database import session
from app.models.entities import ACSProfileObservation, District, Source, Statistic

SOURCE_ID = "acs-2024-5yr"
PERIOD = "2020–2024"
FILES = {
    "house": "acs2024_5yr_house_hi.xlsx",
    "senate": "acs2024_5yr_senate_hi.xlsx",
    "county": "acs2024_5yr_counties_hi.xlsx",
    "state": "acs2024_5yr_state_hi.xlsx",
}
COUNTIES = {
    "Hawaii County, Hawaii": "county-15001",
    "Honolulu County, Hawaii": "county-15003",
    "Kalawao County, Hawaii": "county-15005",
    "Kauai County, Hawaii": "county-15007",
    "Maui County, Hawaii": "county-15009",
}
HEADERS = ("Estimate", "Margin of Error", "Percent", "Percent Margin of Error")


def reported_number(value):
    if value is None:
        return None
    text = str(value).strip().replace(",", "").removeprefix("±").removesuffix("%")
    # Preserve top-coding such as '4,000+' as a marker, not a false exact estimate.
    if not re.fullmatch(r"\d+(?:\.\d+)?", text):
        return None
    return float(text)


def geography_id(label, kind):
    if kind in ("house", "senate"):
        match = re.fullmatch(
            rf"State {'House' if kind == 'house' else 'Senate'} District (\d+) \(2024\); Hawaii", str(label)
        )
        if not match:
            raise ValueError(f"Unrecognized {kind} geography: {label}")
        return f"{kind}-{int(match[1]):02}"
    if kind == "county":
        if label not in COUNTIES:
            raise ValueError(f"Unknown county label: {label}")
        return COUNTIES[label]
    if label != "Hawaii":
        raise ValueError("State workbook must contain only Hawaiʻi")
    return "state-15"


def parse_workbook(path: Path, kind: str):
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    observations = []
    aliases = []
    expected = (
        {f"{kind}-{n:02}" for n in range(1, 52 if kind == "house" else 26)}
        if kind in ("house", "senate")
        else set(COUNTIES.values())
        if kind == "county"
        else {"state-15"}
    )
    try:
        if set(workbook.sheetnames) != {"DP02", "DP03", "DP04", "DP05"}:
            raise ValueError(f"Unexpected profile worksheets: {workbook.sheetnames}")
        for sheet in workbook:
            rows = list(sheet.values)
            if (sheet.max_column - 1) % 4:
                raise ValueError("Each geography requires exactly four measure columns")
            groups = {}
            for offset in range(1, sheet.max_column, 4):
                label = rows[0][offset]
                gid = geography_id(label, kind)
                if gid in groups:
                    raise ValueError("Duplicate workbook geography")
                if tuple(rows[1][offset : offset + 4]) != HEADERS:
                    raise ValueError("Estimate/MOE column headings changed")
                groups[gid] = (offset, str(label))
            if set(groups) != expected:
                raise ValueError(f"Incomplete {kind} geography set")
            # Match definitions, never positions alone.
            for spec in PROFILE_INDICATORS.values():
                if (
                    spec["sheet"] == sheet.title
                    and str(rows[spec["row"] - 1][0]).strip() != spec["expected_label"]
                ):
                    raise ValueError(f"{sheet.title}!A{spec['row']} definition changed")
            section = ""
            for row_number, row in enumerate(rows[2:], 3):
                if row[0] is None:
                    continue
                definition = str(row[0]).strip()
                if not any(v is not None for v in row[1:]):
                    section = definition
                    continue
                for gid, (offset, label) in groups.items():
                    raw = dict(zip(("estimate", "moe", "percent", "percent_moe"), row[offset : offset + 4]))
                    cells = {
                        name: f"{get_column_letter(offset + i + 1)}{row_number}" for i, name in enumerate(raw)
                    }
                    observations.append(
                        {
                            "id": f"profile2024-{gid}-{sheet.title}-{row_number}",
                            "source_id": SOURCE_ID,
                            "geography_id": gid,
                            "profile": sheet.title,
                            "row_number": row_number,
                            "definition": definition,
                            "section": section,
                            "reference_period": PERIOD,
                            "raw_values": raw,
                            **{name: reported_number(value) for name, value in raw.items()},
                            "provenance": {
                                "workbook": path.name,
                                "workbook_sha256": digest,
                                "sheet": sheet.title,
                                "geographic_label": label,
                                "cells": cells,
                                "geographic_vintage": 2024,
                                "workbook_url": f"/v1/sources/{SOURCE_ID}/workbooks/{kind}",
                            },
                        }
                    )
            for code, spec in PROFILE_INDICATORS.items():
                if spec["sheet"] != sheet.title:
                    continue
                row_number = spec["row"]
                for gid, (offset, label) in groups.items():
                    c = spec["column"]
                    raw_value = rows[row_number - 1][offset + c] if c is not None else None
                    raw_moe = rows[row_number - 1][offset + c + 1] if c is not None else None
                    estimate = reported_number(raw_value)
                    moe = reported_number(raw_moe)
                    cell = f"{get_column_letter(offset + c + 1)}{row_number}" if c is not None else None
                    mcell = f"{get_column_letter(offset + c + 2)}{row_number}" if c is not None else None
                    aliases.append(
                        {
                            "id": f"acs2024-{gid}-{code}",
                            "geography_id": gid,
                            "geographic_level": kind,
                            "boundary_vintage": 2022 if kind in ("house", "senate") else None,
                            "indicator": code,
                            "estimate": estimate,
                            "moe": moe,
                            "unit": spec["unit"],
                            "reference_period": PERIOD,
                            "variables": [f"{sheet.title}!{cell}"] if cell else [],
                            "method": "direct_profile_workbook",
                            "status": "available" if estimate is not None else "not_published_or_suppressed",
                            "source_id": SOURCE_ID,
                            "ingested_at": datetime.now(timezone.utc),
                            "provenance": {
                                "source_url": "https://www.census.gov/programs-surveys/acs/data.html",
                                "workbook": path.name,
                                "workbook_sha256": digest,
                                "sheet": sheet.title,
                                "estimate_cell": cell,
                                "moe_cell": mcell,
                                "raw_estimate": raw_value,
                                "raw_moe": raw_moe,
                                "variable_definition": spec["label"],
                                "row_definition": spec["expected_label"],
                                "display_group": spec["display_group"],
                                "display_order": row_number,
                                "geographic_label": label,
                                "acs_geography_vintage": 2024,
                                "boundary_plan": 2022 if kind in ("house", "senate") else None,
                                "moe_confidence": 0.90,
                                "moe_method": "As reported in workbook; no estimated margins added",
                                "workbook_url": f"/v1/sources/{SOURCE_ID}/workbooks/{kind}",
                                "availability_note": "Reported value retained; nonnumeric markers remain unavailable. DP03 poverty counts and universe are not published.",
                            },
                        }
                    )
        return {
            "kind": kind,
            "filename": path.name,
            "sha256": digest,
            "geographies": sorted(expected),
            "observations": observations,
            "statistics": aliases,
        }
    finally:
        workbook.close()


def upsert(db, model, values):
    for start in range(0, len(values), 500):
        statement = insert(model).values(values[start : start + 500])
        db.execute(
            statement.on_conflict_do_update(
                index_elements=["id"],
                set_={
                    c.name: getattr(statement.excluded, c.name)
                    for c in model.__table__.columns
                    if c.name != "id"
                },
            )
        )


def ingest_local(directory=None):
    directory = Path(directory) if directory else ROOT / "data/raw/acs"
    missing = [str(directory / name) for name in FILES.values() if not (directory / name).is_file()]
    if missing:
        raise ValueError("Place the four official workbooks at: " + ", ".join(missing))
    compatibility = json.loads((ROOT / "data/manifests/acs-compatibility.json").read_text())
    if compatibility.get("status") != "reviewed" or compatibility.get("boundary_vintage") != 2022:
        raise ValueError("Review the existing 2024 ACS/2022 plan compatibility manifest first")
    datasets = [parse_workbook(directory / name, kind) for kind, name in FILES.items()]
    with session() as db, db.begin():
        districts = list(db.scalars(select(District)))
        actual = {d.id for d in districts if d.vintage == 2022}
        expected = {
            gid for data in datasets if data["kind"] in ("house", "senate") for gid in data["geographies"]
        }
        if actual != expected or len(districts) != 76:
            raise ValueError(
                "All 76 official 2022 district identifiers must match before statistics can be imported"
            )
        db.merge(
            Source(
                id=SOURCE_ID,
                title="2024 ACS five-year DP02–DP05 profiles (2020–2024)",
                url="https://www.census.gov/programs-surveys/acs/data.html",
                status="approved",
                retrieved_at=datetime.now(timezone.utc),
                metadata_record={
                    "reference_period": PERIOD,
                    "mode": "local_workbooks",
                    "boundary_compatibility": compatibility,
                    "origin": "User-supplied official workbooks; originals retained without modification",
                    "workbooks": [
                        {k: d[k] for k in ("kind", "filename", "sha256", "geographies")} for d in datasets
                    ],
                },
            )
        )
        db.flush()
        for data in datasets:
            upsert(db, ACSProfileObservation, data["observations"])
            upsert(db, Statistic, data["statistics"])
    report = {
        "mode": "local_workbooks",
        "geographies": {d["kind"]: len(d["geographies"]) for d in datasets},
        "statistics_by_geography": {d["kind"]: len(d["statistics"]) for d in datasets},
        "profile_observations_by_geography": {d["kind"]: len(d["observations"]) for d in datasets},
        "mapped_indicators": len(PROFILE_INDICATORS),
        "available_statistic_values": sum(
            s["estimate"] is not None for d in datasets for s in d["statistics"]
        ),
        "unavailable_statistic_values": sum(s["estimate"] is None for d in datasets for s in d["statistics"]),
        "ingested_at": datetime.now(timezone.utc).isoformat(),
    }
    output = ROOT / "data/processed/acs-workbook-ingestion.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, indent=2))
    return report
